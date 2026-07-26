from decimal import Decimal
from unittest.mock import patch

from django.test import TestCase
from openpyxl import Workbook

from apps.reports.services.teacher_workload_excel import (
    TeacherWorkloadExcelService,
)
from apps.staff.models import (
    StaffEmploymentAcademicYear,
)
from apps.staff.tests.factories import (
    create_academic_year,
    create_employment,
)
from apps.reports.exceptions import ReportDataError


class TeacherWorkloadExcelRateTests(TestCase):
    def setUp(self):
        self.academic_year = create_academic_year()

        self.employment = create_employment(
            rate=Decimal("1.00"),
        )

        self.year_record = (
            StaffEmploymentAcademicYear
            .objects
            .create(
                staff_employment=self.employment,
                academic_year=self.academic_year,
                rate=Decimal("0.75"),
            )
        )

    def test_format_rate(self):
        self.assertEqual(
            TeacherWorkloadExcelService.format_rate(
                Decimal("1.00")
            ),
            "1",
        )
        self.assertEqual(
            TeacherWorkloadExcelService.format_rate(
                Decimal("0.75")
            ),
            "0.75",
        )

    def test_placeholder_st_is_replaced_in_y2(self):
        workbook = Workbook()
        worksheet = workbook.active

        worksheet["Y2"] = "Ставка: {st}"

        (
            TeacherWorkloadExcelService
            .replace_placeholders(
                worksheet,
                {
                    "{st}": (
                        TeacherWorkloadExcelService
                        .format_rate(
                            self.year_record.rate
                        )
                    ),
                },
            )
        )

        self.assertEqual(
            worksheet["Y2"].value,
            "Ставка: 0.75",
        )

    def test_year_record_rate_differs_from_current_rate(
        self,
    ):
        self.assertEqual(
            self.employment.rate,
            Decimal("1.00"),
        )
        self.assertEqual(
            self.year_record.rate,
            Decimal("0.75"),
        )

    def test_build_uses_academic_year_record_rate(self):
        result = TeacherWorkloadExcelService.build(
            staff_employment_id=self.employment.pk,
            academic_year=self.academic_year,
        )

        workbook = load_workbook(result)
        worksheet = workbook.active

        self.assertEqual(
            worksheet["Y2"].value,
            "Ставка: 0.75",
        )

    def test_build_rejects_missing_year_record(self):
        self.year_record.delete()

        with self.assertRaisesMessage(
                ReportDataError,
                "не заполнены кадровые данные",
        ):
            TeacherWorkloadExcelService.build(
                staff_employment_id=(
                    self.employment.pk
                ),
                academic_year=self.academic_year,
            )