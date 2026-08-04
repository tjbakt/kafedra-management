from io import BytesIO

import factory
from django.core.files.uploadedfile import (
    SimpleUploadedFile,
)
from openpyxl import Workbook

from apps.reports.models import (
    ExcelReportTemplate,
)
from tests.factories.accounts import UserFactory
from tests.factories.organizations import (
    UniversityFactory,
)


def create_xlsx_file(
    *,
    filename="template.xlsx",
    sheet_name="Отчёт",
    max_columns=25,
    max_rows=6,
    placeholders=None,
):
    """
    Создаёт минимальный корректный XLSX-файл
    для тестов отчётов.
    """
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name

    placeholders = placeholders or {}

    for cell, value in placeholders.items():
        worksheet[cell] = value

    for row in range(1, max_rows + 1):
        for column in range(
            1,
            max_columns + 1,
        ):
            cell = worksheet.cell(
                row=row,
                column=column,
            )

            if cell.value is None:
                cell.value = ""

    output = BytesIO()
    workbook.save(output)
    workbook.close()

    output.seek(0)

    return SimpleUploadedFile(
        filename,
        output.read(),
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
    )


class ExcelReportTemplateFactory(
    factory.django.DjangoModelFactory
):
    class Meta:
        model = ExcelReportTemplate

    university = factory.SubFactory(
        UniversityFactory
    )

    template_type = (
        ExcelReportTemplate
        .Type
        .TEACHER_WORKLOAD
    )

    name = factory.Sequence(
        lambda number: (
            f"Тестовый шаблон {number}"
        )
    )
    version = 1

    file = factory.LazyFunction(
        lambda: create_xlsx_file(
            placeholders={
                "A1": "{FIO}",
                "A2": "{year}",
                "A3": "{st}",
                "A4": "{degree}",
                "B4": "{title}",
            },
        )
    )

    sheet_name = "Отчёт"
    is_active = True
    description = ""

    created_by = factory.SubFactory(
        UserFactory
    )
    updated_by = factory.SelfAttribute(
        "created_by"
    )

    @classmethod
    def global_template(
        cls,
        **kwargs,
    ):
        return cls(
            university=None,
            **kwargs,
        )

    @classmethod
    def department_template(
        cls,
        **kwargs,
    ):
        return cls(
            template_type=(
                ExcelReportTemplate
                .Type
                .DEPARTMENT_WORKLOAD
            ),
            file=create_xlsx_file(
                max_columns=22,
                max_rows=4,
                placeholders={
                    "A1": "{department}",
                    "A2": "{year}",
                },
            ),
            **kwargs,
        )